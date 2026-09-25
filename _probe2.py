import os, unicodedata
import openpyxl

def norm(s):
    return ''.join(c for c in unicodedata.normalize('NFKD', s or '') if c.isascii()).lower()

xlsx = 'C:/Users/thomasphiliippesavar/OneDrive/Documents/systeme_convolutif_spectral_general.xlsx'
if not os.path.exists(xlsx):
    open('_probe_out.txt','w',encoding='utf-8').write('NOT FOUND: '+xlsx)
    raise SystemExit('NOT FOUND')

L = []
def w(s=''): L.append(s)
wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
w('FILE: '+xlsx)
w('SIZE: %s octets' % os.path.getsize(xlsx))
w('SHEETS: '+str(wb.sheetnames))
w('')

def dump(ws, cap):
    i = 0
    for r in ws.iter_rows(values_only=True):
        if cap is not None and i >= cap: break
        cells = [(str(c).replace('\n',' ') if c is not None else '') for c in r]
        w(' | '.join(cells))
        i += 1

for sn in wb.sheetnames:
    n = norm(sn)
    if 'accueil' in n or 'parametre' in n or 'exemple 1-3' in n or 'validation 1-14' in n or 'reference' in n:
        cap = None
    elif 'validation hol' in n:
        cap = 60
    elif 'tests' in n:
        cap = 100
    elif 'checklist' in n:
        cap = 40
    else:
        cap = 30
    try:
        ws = wb[sn]
        w('##### %s | rows=%s cols=%s | mode=%s #####' % (sn, ws.max_row, ws.max_column, ('FULL' if cap is None else 'cap='+str(cap))))
        dump(ws, cap)
    except Exception as e:
        w('ERR %s: %r' % (sn, e))
    w('')

open('_probe_out.txt','w',encoding='utf-8').write('\n'.join(L)+'\n')
print('WRITTEN %d lines' % len(L))
