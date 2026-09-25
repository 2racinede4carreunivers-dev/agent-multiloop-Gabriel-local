import os, unicodedata
import openpyxl

def norm(s):
    return ''.join(c for c in unicodedata.normalize('NFKD', s or '') if c.isascii()).lower()

xlsx = 'C:/Users/thomasphiliippesavar/OneDrive/Documents/systeme_convolutif_spectral_general.xlsx'
wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)

TARGETS = ['code python', 'test 1-14', 'moteur universel', 'systeme general',
           'moteur geo convolutif', 'parametres', 'convolution',
           'validation convolutive', 'suites geo a-b', 'tests digamma',
           'catalogue premiers']
CAP = 600
for sn in wb.sheetnames:
    n = norm(sn)
    match = any(t in n for t in TARGETS)
    cap = None if match and 'catalogue' not in n else (60 if 'catalogue' in n else None)
    if not match:
        continue
    safe = norm(sn).replace(' ', '_')[:20]
    fn = '_sheet_' + safe + '.txt'
    ws = wb[sn]
    out = ['##### %s | rows=%s cols=%s #####' % (sn, ws.max_row, ws.max_column)]
    i = 0
    for r in ws.iter_rows(values_only=True):
        if cap is not None and i >= cap: break
        cells = [(str(c).replace('\n',' ') if c is not None else '') for c in r]
        out.append(' | '.join(cells))
        i += 1
    open(fn, 'w', encoding='utf-8').write('\n'.join(out)+'\n')
    print('DUMP', fn, ':', i, 'lignes')
