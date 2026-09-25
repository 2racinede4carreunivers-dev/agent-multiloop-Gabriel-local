import os, sys
p = 'C:/Users/thomasphiliippesavar/OneDrive/Documents/systeme_convolutif_spectral_general.xlsx'
if not os.path.exists(p):
    sys.exit('NOT FOUND: ' + p)
print('size:', os.path.getsize(p))
import openpyxl
wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
print('SHEETS:', wb.sheetnames)
for sn in wb.sheetnames:
    ws = wb[sn]
    print('=== %s | rows=%s cols=%s ===' % (sn, ws.max_row, ws.max_column))
    cnt = 0
    for r in ws.iter_rows(values_only=True):
        print('  ', r)
        cnt += 1
        if cnt >= 6:
            break
